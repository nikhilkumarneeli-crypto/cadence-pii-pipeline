"""Unit and integration tests for Member 1 components.

Tests:
- src.common.schema
- src.extractors.docx_extractor
- src.router
- src.pipeline
"""

import tempfile
import unittest
from pathlib import Path

import docx

from src.common.schema import Document, PIIEntity, PIIFinding, PipelineResult, Redaction, TextBlock
from src.extractors.docx_extractor import DocxExtractionError, DocxExtractor, extract_docx
from src.pipeline import (
    InvalidComponentError,
    PIIPipeline,
    PipelineExecutionError,
    run_pipeline,
)
from src.router import (
    DocumentRouter,
    ExtractorNotAvailableError,
    UnsupportedFileTypeError,
    detect_file_type,
    get_extractor,
    route_and_extract,
)


class TestMember1Schema(unittest.TestCase):
    """Tests for common schema dataclasses."""

    def test_text_block_creation(self):
        block = TextBlock(
            block_id="b_1",
            text="Fake text content for testing",
            page_number=1,
            metadata={"style": "Normal"},
        )
        self.assertEqual(block.block_id, "b_1")
        self.assertEqual(block.text, "Fake text content for testing")
        self.assertEqual(block.page_number, 1)
        self.assertEqual(block.metadata["style"], "Normal")

    def test_document_full_text_and_get_text(self):
        b1 = TextBlock(block_id="b1", text="First fake paragraph.")
        b2 = TextBlock(block_id="b2", text="Second fake paragraph.")
        doc = Document(
            document_id="doc_01",
            filename="sample.docx",
            file_type="docx",
            blocks=[b1, b2],
            metadata={"author": "Test Author"},
        )
        self.assertEqual(doc.get_text(separator="\n"), "First fake paragraph.\nSecond fake paragraph.")
        self.assertEqual(doc.full_text, "First fake paragraph.\n\nSecond fake paragraph.")

    def test_pii_entity_and_redaction(self):
        entity = PIIEntity(
            entity_type="EMAIL_ADDRESS",
            text="john@example.com",
            start=10,
            end=26,
            confidence=0.95,
            block_id="b1",
        )
        self.assertEqual(entity.entity_type, "EMAIL_ADDRESS")
        self.assertEqual(entity.confidence, 0.95)

        redaction = Redaction(
            entity_type="EMAIL_ADDRESS",
            original_text="john@example.com",
            replacement="<EMAIL_ADDRESS>",
            start=10,
            end=26,
            block_id="b1",
        )
        self.assertEqual(redaction.replacement, "<EMAIL_ADDRESS>")

    def test_pipeline_result(self):
        doc = Document(document_id="d1", filename="d.docx", file_type="docx")
        res = PipelineResult(document=doc, detected_entities=[], redactions=[])
        self.assertEqual(res.document.document_id, "d1")
        self.assertEqual(res.document_id, "d1")
        self.assertEqual(len(res.detected_entities), 0)
        self.assertEqual(len(res.entities), 0)
        self.assertIsNone(res.redacted_text)

        # Member 4 shorthand initialization (document_id, entities)
        res2 = PipelineResult(document_id="d2", entities=[])
        self.assertEqual(res2.document_id, "d2")
        self.assertEqual(res2.detected_entities, [])

    def test_pii_finding(self):
        finding = PIIFinding(
            entity_type="SSN",
            text="000-00-0000",
            start=0,
            end=11,
            confidence=0.85,
            detector="Presidio",
            source="doc_test",
        )
        self.assertEqual(finding.entity_type, "SSN")
        self.assertEqual(finding.confidence, 0.85)
        d = finding.to_dict()
        self.assertEqual(d["detector"], "Presidio")
        self.assertEqual(d["source"], "doc_test")


class TestDocxExtractor(unittest.TestCase):
    """Tests for DOCX extractor."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_docx_with_several_paragraphs(self):
        doc_path = self.dir_path / "paragraphs.docx"
        doc = docx.Document()
        doc.add_heading("Section 1: Overview", level=1)
        doc.add_paragraph("This is the first body paragraph.")
        doc.add_paragraph("This is the second body paragraph.")
        doc.add_heading("Section 2: Details", level=2)
        doc.add_paragraph("Concluding notes.")
        doc.save(doc_path)

        extracted = extract_docx(doc_path)
        self.assertEqual(extracted.file_type, "docx")
        self.assertEqual(len(extracted.blocks), 5)
        self.assertEqual(extracted.blocks[0].text, "Section 1: Overview")
        self.assertTrue(extracted.blocks[0].metadata.get("is_heading"))
        self.assertEqual(extracted.blocks[0].metadata.get("heading_level"), 1)
        self.assertEqual(extracted.blocks[1].text, "This is the first body paragraph.")
        self.assertEqual(extracted.blocks[3].metadata.get("heading_level"), 2)

    def test_docx_with_empty_paragraphs(self):
        doc_path = self.dir_path / "empty_paras.docx"
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1")
        doc.add_paragraph("")  # empty
        doc.add_paragraph("   \t  \n ")  # whitespace
        doc.add_paragraph("Paragraph 2")
        doc.save(doc_path)

        extracted = extract_docx(doc_path)
        self.assertEqual(len(extracted.blocks), 2)
        self.assertEqual(extracted.blocks[0].text, "Paragraph 1")
        self.assertEqual(extracted.blocks[1].text, "Paragraph 2")

    def test_docx_containing_a_table_and_order(self):
        doc_path = self.dir_path / "table.docx"
        doc = docx.Document()
        doc.add_paragraph("Header before table")
        tbl = doc.add_table(rows=2, cols=2)
        tbl.cell(0, 0).text = "Row 0 Col 0"
        tbl.cell(0, 1).text = "Row 0 Col 1"
        tbl.cell(1, 0).text = "Row 1 Col 0"
        tbl.cell(1, 1).text = "Row 1 Col 1"
        doc.add_paragraph("Footer after table")
        doc.save(doc_path)

        extracted = extract_docx(doc_path)
        self.assertEqual(len(extracted.blocks), 6)
        self.assertEqual(extracted.blocks[0].text, "Header before table")
        self.assertEqual(extracted.blocks[0].metadata["block_type"], "paragraph")
        self.assertEqual(extracted.blocks[1].text, "Row 0 Col 0")
        self.assertEqual(extracted.blocks[1].metadata["block_type"], "table_cell")
        self.assertEqual(extracted.blocks[1].metadata["row_index"], 0)
        self.assertEqual(extracted.blocks[1].metadata["col_index"], 0)
        self.assertEqual(extracted.blocks[4].text, "Row 1 Col 1")
        self.assertEqual(extracted.blocks[5].text, "Footer after table")

    def test_docx_containing_fake_pii_preservation(self):
        """Verify that the extractor NEVER redacts or alters fake PII."""
        doc_path = self.dir_path / "pii_report.docx"
        doc = docx.Document()
        doc.add_heading("Employee Report", level=1)
        doc.add_paragraph("Name: John Example")
        doc.add_paragraph("Email: john@example.com")
        doc.add_paragraph("Phone: +1-555-0100")
        doc.add_paragraph("Address: 123 Example Street")
        doc.save(doc_path)

        extracted = extract_docx(doc_path)
        full_text = extracted.full_text

        self.assertIn("John Example", full_text)
        self.assertIn("john@example.com", full_text)
        self.assertIn("+1-555-0100", full_text)
        self.assertIn("123 Example Street", full_text)

        self.assertNotIn("[PERSON]", full_text)
        self.assertNotIn("[EMAIL]", full_text)
        self.assertNotIn("[REDACTED]", full_text)

    def test_existing_member1_docx_fixture(self):
        fixture_path = Path("tests/data/member1_docx/sample_fake_report.docx")
        if fixture_path.exists():
            extracted = extract_docx(fixture_path)
            self.assertEqual(extracted.file_type, "docx")
            self.assertIn("John Example", extracted.full_text)
            self.assertIn("john@example.com", extracted.full_text)
            self.assertIn("+1-555-0100", extracted.full_text)
            self.assertIn("123 Example Street", extracted.full_text)
            self.assertIn("Quality Assurance", extracted.full_text)

    def test_unsupported_extension_docx_extractor(self):
        txt_path = self.dir_path / "note.txt"
        txt_path.write_text("Hello")
        with self.assertRaises(ValueError):
            extract_docx(txt_path)

    def test_missing_file_docx_extractor(self):
        with self.assertRaises(FileNotFoundError):
            extract_docx(self.dir_path / "does_not_exist.docx")

    def test_empty_document(self):
        doc_path = self.dir_path / "empty.docx"
        doc = docx.Document()
        doc.save(doc_path)

        extracted = extract_docx(doc_path)
        self.assertEqual(len(extracted.blocks), 0)
        self.assertEqual(extracted.get_text(), "")

    def test_corrupt_docx(self):
        corrupt_path = self.dir_path / "corrupt.docx"
        corrupt_path.write_bytes(b"PK\x03\x04This is not a real zip or docx")
        with self.assertRaises(DocxExtractionError):
            extract_docx(corrupt_path)


class TestRouter(unittest.TestCase):
    """Tests for format router."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_case_insensitive_extension_handling(self):
        self.assertEqual(detect_file_type("document.DOCX"), "docx")
        self.assertEqual(detect_file_type("document.Docx"), "docx")
        self.assertEqual(detect_file_type("report.PDF"), "pdf")
        self.assertEqual(detect_file_type("report.Pdf"), "pdf")
        self.assertEqual(detect_file_type("slides.PPTX"), "pptx")
        self.assertEqual(detect_file_type("slides.Pptx"), "pptx")

    def test_router_unsupported_extension(self):
        with self.assertRaises(UnsupportedFileTypeError):
            detect_file_type("data.csv")
        with self.assertRaises(UnsupportedFileTypeError):
            detect_file_type("no_extension_file")

    def test_router_docx_dispatch(self):
        doc_path = self.dir_path / "routed.docx"
        doc = docx.Document()
        doc.add_paragraph("Routed Docx Content")
        doc.save(doc_path)

        document = route_and_extract(doc_path)
        self.assertEqual(document.file_type, "docx")
        self.assertEqual(document.blocks[0].text, "Routed Docx Content")

    def test_router_pdf_and_pptx_extractor_available(self):
        # Once Member 2 and Member 3 extractors are integrated, get_extractor returns callables
        pdf_extractor = get_extractor("sample.pdf")
        self.assertTrue(callable(pdf_extractor))
        pptx_extractor = get_extractor("sample.pptx")
        self.assertTrue(callable(pptx_extractor))

    def test_router_pdf_and_pptx_extractor_not_available(self):
        import unittest.mock as mock
        with mock.patch.dict("sys.modules", {"src.extractors.pdf_extractor": None}):
            with self.assertRaises(ExtractorNotAvailableError):
                DocumentRouter().get_extractor("sample.pdf")
        with mock.patch.dict("sys.modules", {"src.extractors.pptx_extractor": None}):
            with self.assertRaises(ExtractorNotAvailableError):
                DocumentRouter().get_extractor("sample.pptx")

    def test_router_custom_extractor_registration(self):
        custom_router = DocumentRouter()
        mock_doc = Document(document_id="mock_pdf", filename="test.pdf", file_type="pdf")
        custom_router.register_extractor("pdf", lambda p, **kw: mock_doc)

        result = custom_router.route_and_extract(Path("test.pdf"))
        self.assertEqual(result.document_id, "mock_pdf")


class TestPipeline(unittest.TestCase):
    """Tests for end-to-end pipeline."""

    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.dir_path = Path(self.temp_dir.name)

        self.test_docx_path = self.dir_path / "employee_summary.docx"
        d = docx.Document()
        d.add_heading("Employee Summary", level=1)
        d.add_paragraph("Employee Name: John Example")
        d.add_paragraph("Contact: john@example.com")
        d.save(self.test_docx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_pipeline_extraction_only(self):
        pipeline = PIIPipeline()
        result = pipeline.process(self.test_docx_path)

        self.assertIsInstance(result, PipelineResult)
        self.assertEqual(result.document.file_type, "docx")
        self.assertEqual(len(result.document.blocks), 3)
        self.assertEqual(len(result.detected_entities), 0)
        self.assertEqual(len(result.redactions), 0)
        self.assertIsNone(result.redacted_text)
        self.assertEqual(result.metadata["stages_executed"], ["extraction"])
        self.assertFalse(result.metadata["is_empty"])

    def test_pipeline_missing_file(self):
        pipeline = PIIPipeline()
        with self.assertRaises(FileNotFoundError):
            pipeline.process(self.dir_path / "nonexistent.docx")

    def test_pipeline_unsupported_file(self):
        bad_file = self.dir_path / "image.png"
        bad_file.write_bytes(b"fake png")
        pipeline = PIIPipeline()
        with self.assertRaises(UnsupportedFileTypeError):
            pipeline.process(bad_file)

    def test_pipeline_empty_document(self):
        empty_doc = self.dir_path / "empty.docx"
        d = docx.Document()
        d.save(empty_doc)

        pipeline = PIIPipeline()
        result = pipeline.process(empty_doc)
        self.assertTrue(result.metadata["is_empty"])
        self.assertEqual(len(result.document.blocks), 0)

    def test_pipeline_with_mock_detector(self):
        class MockDetector:
            def detect(self, doc: Document):
                entities = []
                for block in doc.blocks:
                    if "john@example.com" in block.text:
                        start = block.text.find("john@example.com")
                        entities.append(
                            PIIEntity(
                                entity_type="EMAIL_ADDRESS",
                                text="john@example.com",
                                start=start,
                                end=start + len("john@example.com"),
                                confidence=0.99,
                                block_id=block.block_id,
                            )
                        )
                return entities

        pipeline = PIIPipeline(default_detector=MockDetector())
        result = pipeline.process(self.test_docx_path)

        self.assertIn("detection", result.metadata["stages_executed"])
        self.assertEqual(len(result.detected_entities), 1)
        self.assertEqual(result.detected_entities[0].entity_type, "EMAIL_ADDRESS")
        self.assertEqual(result.detected_entities[0].text, "john@example.com")

    def test_pipeline_with_mock_detector_and_redactor(self):
        def mock_detector(doc: Document):
            return [
                PIIEntity(
                    entity_type="EMAIL_ADDRESS",
                    text="john@example.com",
                    start=9,
                    end=25,
                    block_id="para_2",
                )
            ]

        def mock_redactor(doc: Document, entities):
            redactions = []
            text = doc.full_text
            for ent in entities:
                redactions.append(
                    Redaction(
                        entity_type=ent.entity_type,
                        original_text=ent.text,
                        replacement="<EMAIL_REDACTED>",
                        start=ent.start,
                        end=ent.end,
                        block_id=ent.block_id,
                    )
                )
                text = text.replace(ent.text, "<EMAIL_REDACTED>")
            return redactions, text

        pipeline = PIIPipeline()
        result = pipeline.process(
            self.test_docx_path,
            detector=mock_detector,
            redactor=mock_redactor,
        )

        self.assertEqual(
            result.metadata["stages_executed"],
            ["extraction", "detection", "redaction"],
        )
        self.assertEqual(len(result.detected_entities), 1)
        self.assertEqual(len(result.redactions), 1)
        self.assertEqual(result.redactions[0].replacement, "<EMAIL_REDACTED>")
        self.assertIn("<EMAIL_REDACTED>", result.redacted_text)
        self.assertNotIn("john@example.com", result.redacted_text)

    def test_pipeline_invalid_detector_redactor(self):
        pipeline = PIIPipeline()
        with self.assertRaises(InvalidComponentError):
            pipeline.process(self.test_docx_path, detector=12345)

        with self.assertRaises(InvalidComponentError):
            pipeline.process(self.test_docx_path, redactor="not_a_redactor")

    def test_convenience_run_pipeline(self):
        result = run_pipeline(self.test_docx_path)
        self.assertIsInstance(result, PipelineResult)
        self.assertEqual(result.document.file_type, "docx")

    def test_pipeline_with_text_only_detector(self):
        """Tests integration with Member 4's PIIDetector interface (detect(text: str))."""
        class MockMember4Detector:
            def detect(self, text: str):
                if not isinstance(text, str):
                    raise TypeError("text must be a string")
                findings = []
                if "john@example.com" in text:
                    start = text.index("john@example.com")
                    findings.append(
                        PIIFinding(
                            entity_type="EMAIL",
                            text="john@example.com",
                            start=start,
                            end=start + len("john@example.com"),
                            confidence=0.99,
                            detector="Presidio",
                        )
                    )
                return findings

        pipeline = PIIPipeline()
        result = pipeline.process(self.test_docx_path, detector=MockMember4Detector())
        self.assertEqual(len(result.detected_entities), 1)
        self.assertEqual(result.detected_entities[0].entity_type, "EMAIL")
        self.assertEqual(result.detected_entities[0].block_id, "para_2")

    def test_pipeline_with_member5_4tuple_redactor(self):
        """Tests integration with Member 5's redact_document interface returning 4-tuple."""
        class MockMember5Redactor:
            def redact_document(self, doc: Document, entities):
                redactions = [
                    Redaction(
                        entity_type="EMAIL",
                        original_text="john@example.com",
                        replacement="[EMAIL_1]",
                        start=0,
                        end=16,
                        block_id="para_2",
                    )
                ]
                redacted_text = doc.full_text.replace("john@example.com", "[EMAIL_1]")
                trace_log = [{"tag": "[EMAIL_1]", "pii_type": "EMAIL"}]
                return redactions, redacted_text, doc, trace_log

        pipeline = PIIPipeline()
        result = pipeline.process(
            self.test_docx_path,
            detector=lambda d: [PIIFinding("EMAIL", "john@example.com", 0, 16, 1.0, "test")],
            redactor=MockMember5Redactor(),
        )
        self.assertEqual(len(result.redactions), 1)
        self.assertEqual(result.redactions[0].replacement, "[EMAIL_1]")
        self.assertIn("traceability_log", result.metadata)

    def test_pipeline_end_to_end_docx_with_real_redactor(self):
        """Tests end-to-end processing of a real DOCX document with Member 5's PIIRedactor."""
        from src.redaction.redactor import PIIRedactor
        sample_path = Path("tests/data/member1_docx/sample_fake_report.docx")
        if not sample_path.exists():
            self.skipTest("Sample DOCX not found")

        pipeline = PIIPipeline()
        def detector(doc):
            findings = []
            for block in doc.blocks:
                if "john@example.com" in block.text:
                    idx = block.text.index("john@example.com")
                    findings.append(
                        PIIEntity(
                            entity_type="EMAIL",
                            text="john@example.com",
                            start=idx,
                            end=idx + len("john@example.com"),
                            block_id=block.block_id,
                        )
                    )
            return findings

        result = pipeline.process(sample_path, detector=detector, redactor=PIIRedactor())
        self.assertEqual(result.document.file_type, "docx")
        self.assertEqual(len(result.redactions), 1)
        self.assertIn("[EMAIL_", result.redacted_text)
        self.assertNotIn("john@example.com", result.redacted_text)
        self.assertIn("exposure_report", result.metadata)

    def test_pipeline_end_to_end_pdf_with_real_redactor(self):
        """Tests end-to-end processing of a real PDF document with Member 5's PIIRedactor."""
        from src.redaction.redactor import PIIRedactor
        sample_path = Path("tests/data/sample_digital_policy_pii.pdf")
        if not sample_path.exists():
            self.skipTest("Sample PDF not found")

        pipeline = PIIPipeline()
        def detector(doc):
            findings = []
            target = "vikram.rao@cadencebank.com"
            for block in doc.blocks:
                if target in block.text:
                    idx = block.text.index(target)
                    findings.append(
                        PIIEntity(
                            entity_type="EMAIL",
                            text=target,
                            start=idx,
                            end=idx + len(target),
                            block_id=block.block_id,
                        )
                    )
            return findings

        result = pipeline.process(sample_path, detector=detector, redactor=PIIRedactor())
        self.assertEqual(result.document.file_type, "pdf")
        self.assertEqual(len(result.redactions), 1)
        self.assertIn("[EMAIL_", result.redacted_text)
        self.assertIn("exposure_report", result.metadata)

    def test_pipeline_end_to_end_pptx_with_real_redactor(self):
        """Tests end-to-end processing of a real PPTX document with Member 5's PIIRedactor."""
        from src.redaction.redactor import PIIRedactor
        sample_path = Path("tests/data/member3_pptx/test1_text_boxes_pii.pptx")
        if not sample_path.exists():
            self.skipTest("Sample PPTX not found")

        pipeline = PIIPipeline()
        def detector(doc):
            findings = []
            target = "Eleanor Vance"
            for block in doc.blocks:
                if target in block.text:
                    idx = block.text.index(target)
                    findings.append(
                        PIIEntity(
                            entity_type="PERSON",
                            text=target,
                            start=idx,
                            end=idx + len(target),
                            block_id=block.block_id,
                        )
                    )
            return findings

        result = pipeline.process(sample_path, detector=detector, redactor=PIIRedactor())
        self.assertEqual(result.document.file_type, "pptx")
        self.assertEqual(len(result.redactions), 1)
        self.assertIn("[PERSON_", result.redacted_text)
        self.assertIn("exposure_report", result.metadata)


if __name__ == "__main__":
    unittest.main()
