"""
Cadence PII Shield — Enterprise Compliance & Redaction Engine
=============================================================
Enterprise-grade Cyber Risk & Compliance dashboard for PII detection,
redaction, audit logging, and native-format sanitized exports.
Single-file Streamlit app (app/demo.py).

Pipeline: Ingestion -> Extraction -> Detection -> Redaction -> Audit Report -> Native Export
"""

import csv
import hashlib
import html
import io
import json
import logging
import os
import re
import sys
import tempfile
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

logger = logging.getLogger(__name__)

# Ensure TESSDATA_PREFIX is configured on Windows
for _tp in [
    os.environ.get("TESSDATA_PREFIX"),
    r"C:\Program Files\Tesseract-OCR\tessdata",
    r"C:\Program Files (x86)\Tesseract-OCR\tessdata",
    str(Path.home() / "AppData" / "Local" / "Programs" / "Tesseract-OCR" / "tessdata"),
]:
    if _tp and Path(_tp).exists():
        os.environ["TESSDATA_PREFIX"] = str(_tp)
        break

# --------------------------------------------------------------------------- #
# 0. PATH SETUP
# --------------------------------------------------------------------------- #
APP_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = APP_DIR.parent
for _p in (str(PROJECT_ROOT), str(APP_DIR)):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import streamlit as st

try:
    import pandas as pd
    _HAVE_PANDAS = True
except Exception:
    pd = None
    _HAVE_PANDAS = False


# --------------------------------------------------------------------------- #
# 1. SCHEMA CONTRACTS
# --------------------------------------------------------------------------- #
try:
    from src.common.schema import TextBlock, Document, PIIEntity, Redaction
    _HAVE_SCHEMA = True
except Exception:
    _HAVE_SCHEMA = False

    @dataclass
    class TextBlock:
        block_id: str
        text: str
        page_number: Optional[int] = None
        metadata: Dict[str, Any] = field(default_factory=dict)

    @dataclass
    class Document:
        document_id: str
        filename: str
        file_type: str
        blocks: List[TextBlock] = field(default_factory=list)
        metadata: Dict[str, Any] = field(default_factory=dict)

    @dataclass
    class PIIEntity:
        entity_type: str
        text: str
        start: int
        end: int
        confidence: float = 1.0
        block_id: Optional[str] = None
        metadata: Dict[str, Any] = field(default_factory=dict)

    @dataclass
    class Redaction:
        entity_type: str
        original_text: str
        replacement: str
        start: int
        end: int
        block_id: Optional[str] = None


# --------------------------------------------------------------------------- #
# 2. DOMAIN CONSTANTS
# --------------------------------------------------------------------------- #
RISK_ORDER = {"CRITICAL": 4, "HIGH": 3, "MEDIUM": 2, "LOW": 1, "NONE": 0}

CANONICAL_TYPE_MAP = {
    "PERSON": "PERSON",
    "PERSON_NAME": "PERSON",
    "EMPLOYEE_ID": "EMPLOYEE_ID",
    "EMAIL": "EMAIL",
    "EMAIL_ADDRESS": "EMAIL",
    "PHONE": "PHONE",
    "PHONE_NUMBER": "PHONE",
    "SSN": "SSN",
    "US_SSN": "SSN",
    "GOVERNMENT_ID": "GOVERNMENT_ID",
    "PAN": "GOVERNMENT_ID",
    "AADHAAR": "GOVERNMENT_ID",
    "PASSPORT": "GOVERNMENT_ID",
    "CREDIT_CARD": "CREDIT_CARD",
    "CARD_NUMBER": "CREDIT_CARD",
    "ACCOUNT_NUMBER": "ACCOUNT_NUMBER",
    "DATE_OF_BIRTH": "DATE_OF_BIRTH",
    "DOB": "DATE_OF_BIRTH",
    "ADDRESS": "ADDRESS",
    "HOME_ADDRESS": "ADDRESS",
    "DATE": "DATE",
}

TAG_PREFIX = {
    "PERSON": "PERSON",
    "PERSON_NAME": "PERSON",
    "EMPLOYEE_ID": "EMP_ID",
    "EMAIL": "EMAIL",
    "EMAIL_ADDRESS": "EMAIL",
    "PHONE": "PHONE",
    "PHONE_NUMBER": "PHONE",
    "US_SSN": "SSN",
    "SSN": "SSN",
    "GOVERNMENT_ID": "GOV_ID",
    "PAN": "GOV_ID",
    "AADHAAR": "GOV_ID",
    "PASSPORT": "GOV_ID",
    "CREDIT_CARD": "CC",
    "CARD_NUMBER": "CC",
    "ACCOUNT_NUMBER": "ACCT",
    "DATE_OF_BIRTH": "DOB",
    "DOB": "DOB",
    "ADDRESS": "ADDR",
    "HOME_ADDRESS": "ADDR",
    "DATE": "DATE",
}

TYPE_SHORT = {
    "PERSON": "PERSON", "PERSON_NAME": "PERSON",
    "EMPLOYEE_ID": "EMP_ID",
    "EMAIL_ADDRESS": "EMAIL", "EMAIL": "EMAIL",
    "PHONE_NUMBER": "PHONE", "PHONE": "PHONE",
    "US_SSN": "SSN", "SSN": "SSN",
    "GOVERNMENT_ID": "GOV_ID", "PAN": "PAN", "AADHAAR": "AADHAAR", "PASSPORT": "PASSPORT",
    "CREDIT_CARD": "CREDIT_CARD", "CARD_NUMBER": "CREDIT_CARD",
    "ACCOUNT_NUMBER": "ACCT_NO",
    "DATE_OF_BIRTH": "DOB", "DOB": "DOB",
    "ADDRESS": "ADDRESS", "HOME_ADDRESS": "ADDRESS",
    "DATE": "DATE",
}

SENSITIVITY = {
    "US_SSN": "CRITICAL", "SSN": "CRITICAL", "GOVERNMENT_ID": "CRITICAL",
    "PAN": "CRITICAL", "AADHAAR": "CRITICAL", "PASSPORT": "CRITICAL",
    "CREDIT_CARD": "CRITICAL", "CARD_NUMBER": "CRITICAL", "ACCOUNT_NUMBER": "CRITICAL",
    "EMPLOYEE_ID": "HIGH", "DATE_OF_BIRTH": "HIGH", "DOB": "HIGH",
    "EMAIL_ADDRESS": "MEDIUM", "EMAIL": "MEDIUM",
    "PHONE_NUMBER": "MEDIUM", "PHONE": "MEDIUM",
    "ADDRESS": "MEDIUM", "HOME_ADDRESS": "MEDIUM",
    "PERSON": "LOW", "PERSON_NAME": "LOW",
    "DATE": "LOW",
}

ALL_TYPES = [
    "PERSON",
    "EMPLOYEE_ID",
    "EMAIL",
    "PHONE",
    "SSN",
    "GOVERNMENT_ID",
    "CREDIT_CARD",
    "ACCOUNT_NUMBER",
    "DATE_OF_BIRTH",
    "ADDRESS",
    "DATE",
]
PIPELINE_STAGES = ["Ingestion", "Extraction", "Detection", "Redaction", "Audit Report"]


# --------------------------------------------------------------------------- #
# 3. BUILT-IN FALLBACK EXTRACTORS
# --------------------------------------------------------------------------- #
def _ocr_page_fitz(path: Union[str, Path], page_number: int) -> Tuple[str, float]:
    try:
        import fitz
        doc = fitz.open(str(path))
        page = doc[page_number - 1]
        tess_path = os.environ.get("TESSDATA_PREFIX")
        ocr_kwargs: Dict[str, Any] = {"flags": 0, "language": "eng", "dpi": 200, "full": True}
        if tess_path:
            ocr_kwargs["tessdata"] = tess_path
        tp = page.get_textpage_ocr(**ocr_kwargs)
        text = page.get_text(textpage=tp).strip()
        doc.close()
        return text, (0.62 if text else 0.0)
    except Exception as exc:
        logger.error(f"OCR failed for page {page_number} of {path}: {exc}")
        return "", 0.0


def _extract_pdf(path: Union[str, Path]) -> List[TextBlock]:
    import pdfplumber
    blocks: List[TextBlock] = []
    with pdfplumber.open(str(path)) as pdf:
        for pno, page in enumerate(pdf.pages, start=1):
            words = page.extract_words(use_text_flow=True) or []
            if not words:
                text, conf = _ocr_page_fitz(path, pno)
                if text:
                    blocks.append(TextBlock(
                        block_id=f"p{pno}_ocr1", text=text, page_number=pno,
                        metadata={"block_type": "paragraph", "source_type": "scanned_ocr",
                                  "is_heading": False, "position": None, "confidence": conf,
                                  "low_confidence_flag": conf < 0.75,
                                  "low_confidence_words": []}))
                continue
            clusters: "OrderedDict[int, List[dict]]" = OrderedDict()
            for w in sorted(words, key=lambda x: (x["top"], x["x0"])):
                key = None
                for k in clusters:
                    if abs(k - w["top"]) <= 4:
                        key = k
                        break
                if key is None:
                    key = round(w["top"])
                    clusters[key] = []
                clusters[key].append(w)
            for li, (top, ws) in enumerate(clusters.items()):
                line_text = " ".join(x["text"] for x in sorted(ws, key=lambda x: x["x0"]))
                blocks.append(TextBlock(
                    block_id=f"p{pno}_b{li + 1:03d}", text=line_text, page_number=pno,
                    metadata={"block_type": "paragraph", "source_type": "digital_pdf",
                              "is_heading": False,
                              "position": {"left": min(x["x0"] for x in ws), "top": top,
                                           "width": max(x["x1"] for x in ws) - min(x["x0"] for x in ws),
                                           "height": max(x["height"] for x in ws)},
                              "confidence": 0.98, "low_confidence_flag": False,
                              "low_confidence_words": []}))
    return blocks


def _extract_docx(path: Union[str, Path]) -> List[TextBlock]:
    import docx as _docx_mod
    d = _docx_mod.Document(str(path))
    blocks: List[TextBlock] = []
    i = 0
    for p in d.paragraphs:
        t = (p.text or "").strip()
        if not t:
            continue
        i += 1
        style = (p.style.name or "").lower() if p.style is not None else ""
        is_head = style.startswith("heading") or style == "title"
        blocks.append(TextBlock(
            block_id=f"b{i:04d}", text=t,
            metadata={"block_type": "heading" if is_head else "paragraph",
                      "source_type": "digital", "is_heading": is_head,
                      "confidence": 0.99, "low_confidence_flag": False,
                      "low_confidence_words": []}))
    for tbl in d.tables:
        for r in tbl.rows:
            for c in r.cells:
                t = (c.text or "").strip()
                if not t:
                    continue
                i += 1
                blocks.append(TextBlock(
                    block_id=f"b{i:04d}", text=t,
                    metadata={"block_type": "table_cell", "source_type": "digital",
                              "is_heading": False, "confidence": 0.99,
                              "low_confidence_flag": False, "low_confidence_words": []}))
    return blocks


def _extract_pptx(path: Union[str, Path]) -> List[TextBlock]:
    from pptx import Presentation
    prs = Presentation(str(path))
    blocks: List[TextBlock] = []
    i = 0
    for sno, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                is_title = shape == slide.shapes.title
                for para in shape.text_frame.paragraphs:
                    t = para.text.strip()
                    if not t:
                        continue
                    i += 1
                    is_head = is_title or para.level == 0
                    blocks.append(TextBlock(
                        block_id=f"s{sno}_b{i:03d}", text=t, page_number=sno,
                        metadata={"block_type": "heading" if is_head else "paragraph",
                                  "source_type": "digital", "is_heading": is_head,
                                  "slide_number": sno, "confidence": 0.99,
                                  "low_confidence_flag": False, "low_confidence_words": []}))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        t = (cell.text or "").strip()
                        if not t:
                            continue
                        i += 1
                        blocks.append(TextBlock(
                            block_id=f"s{sno}_b{i:03d}", text=t, page_number=sno,
                            metadata={"block_type": "table_cell", "source_type": "digital",
                                      "is_heading": False, "slide_number": sno,
                                      "confidence": 0.99, "low_confidence_flag": False,
                                      "low_confidence_words": []}))
    return blocks


def _fallback_route_and_extract(path: Union[str, Path]) -> Document:
    p = Path(path)
    ext = p.suffix.lower().lstrip(".")
    if ext == "docx":
        blocks = _extract_docx(p)
    elif ext == "pptx":
        blocks = _extract_pptx(p)
    elif ext == "pdf":
        blocks = _extract_pdf(p)
    else:
        raise ValueError(f"Unsupported file type: {ext}")
    return Document(
        document_id=hashlib.sha256(str(p).encode()).hexdigest()[:12],
        filename=p.name, file_type=ext, blocks=blocks,
        metadata={"extraction_engine": "builtin_fallback", "source_path": str(p)})


def route_and_extract(path: Union[str, Path]) -> Document:
    try:
        from src.router import route_and_extract as _team_route
        doc = _team_route(path)
        if doc is not None and getattr(doc, "blocks", None) is not None:
            return doc
    except Exception:
        pass
    return _fallback_route_and_extract(path)


# --------------------------------------------------------------------------- #
# 4. DETECTION (ENGINE + BUILT-IN HYBRID REGEX)
# --------------------------------------------------------------------------- #
try:
    from app.robust_redactor import (
        is_governance_term,
        expand_name_span,
        GOVERNANCE_STOPWORDS,
        EXECUTIVE_ROSTER,
        is_valid_person_name,
        is_forbidden_name,
        FORBIDDEN_NAME_WORDS,
    )
except ImportError:
    try:
        from robust_redactor import (
            is_governance_term,
            expand_name_span,
            GOVERNANCE_STOPWORDS,
            EXECUTIVE_ROSTER,
            is_valid_person_name,
            is_forbidden_name,
            FORBIDDEN_NAME_WORDS,
        )
    except ImportError:
        from src.redaction.robust_redactor import (
            is_governance_term,
            expand_name_span,
            GOVERNANCE_STOPWORDS,
            EXECUTIVE_ROSTER,
            is_valid_person_name,
            is_forbidden_name,
            FORBIDDEN_NAME_WORDS,
        )

_RE_EMAIL = re.compile(r"\b[A-Za-z0-9._%+-]+(?:\s*@\s*)[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b")
_RE_SSN = re.compile(r"(?<!\d)\d{3}-\d{2}-\d{4}(?!\d)")
_RE_EMP = re.compile(r"\b(?:EMP|DIR)-\d{3,8}\b")
_RE_EMP_LABEL = re.compile(r"(?i)\bemployee\s*(?:id|number|no\.?)\s*[:\-#]?\s*([A-Z]{0,4}-?\d{4,})\b")
_RE_PAN = re.compile(r"\b[A-Z]{5}\d{4}[A-Z]\b")
_RE_PHONE = re.compile(r"(?<!\d)(?:\+?\d{1,3}[\s.\-]?)?\(?\d{2,4}\)?[\s.\-]\d{3,4}[\s.\-]\d{4}(?!\d)")
_RE_DOB = re.compile(
    r"(?i)\b(?:date\s+of\s+birth|d\.?\s?o\.?\s?b\.?|birthdate|born\s+on)\s*[:\-]?\s*"
    r"([A-Za-z]{3,9}\s+\d{1,2},?\s+\d{4}|\d{1,2}[/-]\d{1,2}[/-]\d{2,4})")
_RE_ADDR = re.compile(
    r"\b\d{1,5}\s+[A-Za-z0-9.'-]+(?:\s+[A-Za-z0-9.'-]+)*\s+"
    r"(?:Street|St\b|Avenue|Ave\b|Road|Rd\b|Boulevard|Blvd\b|Lane|Ln\b|Drive|Dr\b|"
    r"Court|Ct\b|Way|Place|Pl\b|Terrace|Ter\b|Parkway|Pkwy\b|Circle|Cir\b|Highway|Hwy\b)"
    r"(?:\s+(?:Apt|Suite|Unit|Ste)\s*\.?\s*\w+)?(?:,\s*[A-Za-z .]+,?\s*[A-Z]{2}\s+\d{5}(?:-\d{4})?)?",
    re.IGNORECASE)
_RE_CC_CANDIDATE = re.compile(r"(?<!\d)(?:\d[ -]?){12,19}(?!\d)")
_RE_NAME_LABEL = re.compile(
    r"(?i:\b(?:reviewer|approver|person\s+name|full\s+name|employee\s+name|contact\s+person|"
    r"prepared\s+by|approved\s+by|submitted\s+by|reviewed\s+by|signatory|attested\s+by|reported\s+by)\b)\s*[:\-]\s*"
)
_RE_NAME_VALUE = re.compile(r"([A-Z][a-zA-Z'.-]+(?:\s+[A-Z]\.?)?\s+[A-Z][a-zA-Z'.-]+)")
_ROLE_LABEL_PREFIXES = (
    r"policy\s+owner|policy\s+author|legal\s+reviewer|privacy\s+reviewer|approval\s+authority|"
    r"chaired\s+by|incumbent|prepared\s+by|approved\s+by|submitted\s+by|reported\s+by|"
    r"reviewed\s+by|attested\s+by|signatory|contact\s+person|lead\s+reviewer|chief\s+reviewer|"
    r"committee\s+chair|chairperson|peer\s+reviewer|qa\s+reviewer"
)
_ROLE_LABEL_PATTERN = re.compile(
    rf"(?i:\b(?:{_ROLE_LABEL_PREFIXES})\b)\s*[:\-]\s*"
    r"([A-Z][a-z]+(?:\s+[A-Z]\.?)*(?:\s+[A-Z][a-z]+(?:-[A-Z][a-z]+)?)+)"
)
_NAME_STOPWORDS = {
    "the", "and", "for", "with", "this", "that", "from", "policy", "report", "review",
    "approval", "approved", "section", "appendix", "control", "controls", "cadence",
    "optiv", "new", "york", "united", "states", "document", "version", "date",
} | FORBIDDEN_NAME_WORDS
_RE_NAME_CELL = re.compile(r"^([A-Z][a-zA-Z'.-]+\s+[A-Z][a-zA-Z'.-]+)$")


def _luhn_ok(digits: str) -> bool:
    total, alt = 0, False
    for ch in reversed(digits):
        d = int(ch)
        if alt:
            d *= 2
            if d > 9:
                d -= 9
        total += d
        alt = not alt
    return total % 10 == 0


def _builtin_detect(document: Document) -> List[PIIEntity]:
    entities: List[PIIEntity] = []
    for block in document.blocks:
        text = block.text or ""
        if not text.strip():
            continue

        def add(etype: str, start: int, end: int, conf: float, detector: str) -> None:
            snippet = text[start:end]
            if not snippet.strip():
                return
            if etype in ("PERSON", "PERSON_NAME") and not is_valid_person_name(snippet):
                return
            entities.append(PIIEntity(
                entity_type=etype, text=snippet, start=start, end=end,
                confidence=conf, block_id=block.block_id,
                metadata={"detector": detector, "sensitivity": SENSITIVITY.get(etype, "LOW"),
                          "source": (block.metadata or {}).get("source_type", "digital")}))

        for m in _RE_EMAIL.finditer(text):
            clean_e = m.group(0).rstrip(".,;:!?)]}>(\"'")
            add("EMAIL_ADDRESS", m.start(), m.start() + len(clean_e), 0.99, "regex:email")
        for m in _RE_SSN.finditer(text):
            add("US_SSN", m.start(), m.end(), 0.97, "regex:ssn_format")
        for m in _RE_EMP.finditer(text):
            add("EMPLOYEE_ID", m.start(), m.end(), 0.99, "regex:emp_id_format")
        for m in _RE_EMP_LABEL.finditer(text):
            add("EMPLOYEE_ID", m.start(1), m.end(1), 0.90, "context:employee_id_label")
        for m in _RE_PAN.finditer(text):
            add("GOVERNMENT_ID", m.start(), m.end(), 0.98, "regex:pan_checksum_format")
        for m in _RE_PHONE.finditer(text):
            add("PHONE_NUMBER", m.start(), m.end(), 0.90, "regex:phone")
        for m in _RE_DOB.finditer(text):
            add("DATE_OF_BIRTH", m.start(1), m.end(1), 0.95, "context:dob_label")
        for m in _RE_ADDR.finditer(text):
            add("ADDRESS", m.start(), m.end(), 0.90, "regex:street_address")
        for m in _RE_CC_CANDIDATE.finditer(text):
            digits = re.sub(r"[ -]", "", m.group(0))
            if len(digits) < 12:
                continue
            if 13 <= len(digits) <= 19 and _luhn_ok(digits):
                add("CREDIT_CARD", m.start(), m.end(), 0.97, "checksum:luhn")
            elif 8 <= len(digits) <= 17:
                add("ACCOUNT_NUMBER", m.start(), m.end(), 0.85, "format:account_number")

        # Executive Roster matching
        for exec_name in EXECUTIVE_ROSTER:
            for m in re.finditer(rf"\b{re.escape(exec_name)}\b", text, re.IGNORECASE):
                if is_valid_person_name(m.group(0)):
                    add("PERSON", m.start(), m.end(), 0.99, "gazetteer:executive_roster")

        # Role & Header context matching
        for m in _ROLE_LABEL_PATTERN.finditer(text):
            cand_name = m.group(1).strip().rstrip(".,;:")
            c_start = m.start(1)
            c_end = c_start + len(cand_name)
            if is_valid_person_name(cand_name):
                est, een, eval_ = expand_name_span(text, c_start, c_end)
                if is_valid_person_name(eval_):
                    add("PERSON", est, een, 0.95, "context:role_label")

        for m in _RE_NAME_LABEL.finditer(text):
            window = text[m.end():m.end() + 80]
            nm = _RE_NAME_VALUE.match(window)
            if nm:
                c_start = m.end() + nm.start(1)
                c_end = m.end() + nm.end(1)
                cand_val = text[c_start:c_end]
                if is_valid_person_name(cand_val):
                    est, een, eval_ = expand_name_span(text, c_start, c_end)
                    if is_valid_person_name(eval_):
                        add("PERSON", est, een, 0.93, "context:name_label")
    return entities


def _normalize_engine_output(result: Any, document: Document) -> Optional[List[PIIEntity]]:
    if result is None:
        return None
    if isinstance(result, dict):
        for key in ("entities", "pii_entities", "findings"):
            if key in result:
                result = result[key]
                break
    if not isinstance(result, (list, tuple)):
        return None
    valid_blocks = {b.block_id for b in document.blocks}
    out: List[PIIEntity] = []
    for item in result:
        try:
            if isinstance(item, PIIEntity):
                e = item
            elif isinstance(item, dict):
                e = PIIEntity(
                    entity_type=item.get("entity_type") or item.get("type") or "PERSON",
                    text=item.get("text", ""),
                    start=int(item.get("start", 0)),
                    end=int(item.get("end", 0)),
                    confidence=float(item.get("confidence", 1.0)),
                    block_id=item.get("block_id"),
                    metadata=item.get("metadata", {}) or {})
            else:
                e = PIIEntity(
                    entity_type=getattr(item, "entity_type", "PERSON"),
                    text=getattr(item, "text", ""),
                    start=int(getattr(item, "start", 0)),
                    end=int(getattr(item, "end", 0)),
                    confidence=float(getattr(item, "confidence", 1.0)),
                    block_id=getattr(item, "block_id", None),
                    metadata=getattr(item, "metadata", {}) or {})
            if e.block_id and e.block_id not in valid_blocks:
                continue
            if not e.metadata.get("sensitivity"):
                e.metadata["sensitivity"] = SENSITIVITY.get(e.entity_type, "LOW")
            if not e.metadata.get("detector"):
                e.metadata["detector"] = "src.detection.engine"
            out.append(e)
        except Exception:
            continue
    return out


def _engine_detect(document: Document) -> Optional[List[PIIEntity]]:
    try:
        from src.detection import engine as eng
    except Exception:
        return None
    for fn_name in ("detect_pii_in_document", "analyze", "detect"):
        fn = getattr(eng, fn_name, None)
        if callable(fn):
            try:
                ents = _normalize_engine_output(fn(document), document)
                if ents is not None:
                    return ents
            except Exception:
                continue
    cls = getattr(eng, "PIIDetector", None)
    if cls is not None:
        try:
            det = cls()
            for m in ("detect", "detect_pii", "analyze"):
                fn = getattr(det, m, None)
                if callable(fn):
                    ents = _normalize_engine_output(fn(document), document)
                    if ents is not None:
                        return ents
        except Exception:
            return None
    return None


def _resolve_overlaps(entities: List[PIIEntity]) -> List[PIIEntity]:
    by_block: Dict[str, List[PIIEntity]] = {}
    for e in entities:
        by_block.setdefault(e.block_id or "", []).append(e)
    kept: List[PIIEntity] = []
    for bids in by_block.values():
        bids.sort(key=lambda e: (e.start, -(e.end - e.start), -e.confidence))
        last_end = -1
        for e in bids:
            if e.start >= last_end:
                kept.append(e)
                last_end = e.end
    return kept


def detect_pii(document: Document, threshold: float) -> List[PIIEntity]:
    eng_ents = _engine_detect(document) or []
    builtin_ents = _builtin_detect(document) or []
    combined = list(eng_ents) + list(builtin_ents)
    filtered = [
        e for e in combined
        if e.confidence >= threshold
        and (e.text or "").strip()
        and not (e.entity_type in ("PERSON", "PERSON_NAME") and not is_valid_person_name(e.text))
    ]
    return _resolve_overlaps(filtered)


# --------------------------------------------------------------------------- #
# 5. REDACTION & AUDIT REPORTING
# --------------------------------------------------------------------------- #
def _fallback_redact_document(document: Document, entities: List[PIIEntity]
                              ) -> Tuple[List[Redaction], str, Document, List[Dict[str, Any]]]:
    counters: Dict[str, int] = {}
    redactions: List[Redaction] = []
    trace: List[Dict[str, Any]] = []
    per_block: Dict[str, List[PIIEntity]] = {}
    for e in entities:
        per_block.setdefault(e.block_id or "", []).append(e)

    new_blocks: List[TextBlock] = []
    for block in document.blocks:
        text = block.text or ""
        block_ents = sorted(per_block.get(block.block_id, []),
                            key=lambda e: e.start, reverse=True)
        for e in block_ents:
            prefix = TAG_PREFIX.get(e.entity_type, e.entity_type.upper())
            counters[prefix] = counters.get(prefix, 0) + 1
            tag = f"[{prefix}_{counters[prefix]:02d}]"
            original = text[e.start:e.end]
            redactions.append(Redaction(entity_type=e.entity_type, original_text=original,
                                        replacement=tag, start=e.start, end=e.end,
                                        block_id=block.block_id))
            trace.append({
                "file": document.filename,
                "page_or_slide": block.page_number,
                "block_id": block.block_id,
                "block_type": (block.metadata or {}).get("block_type", "paragraph"),
                "pii_type": e.entity_type,
                "tag": tag,
                "original_text": original,
                "start": e.start,
                "end": e.end,
                "detector": (e.metadata or {}).get("detector", "unknown"),
                "confidence": round(float(e.confidence), 4),
                "sensitivity": (e.metadata or {}).get("sensitivity", "LOW"),
            })
            text = text[:e.start] + tag + text[e.end:]
        new_blocks.append(TextBlock(
            block_id=block.block_id, text=text, page_number=block.page_number,
            metadata=dict(block.metadata or {})))

    redacted_doc = Document(document_id=document.document_id, filename=document.filename,
                            file_type=document.file_type, blocks=new_blocks,
                            metadata=dict(document.metadata or {}))
    redacted_text = "\n\n".join(nb.text for nb in new_blocks)
    return redactions, redacted_text, redacted_doc, trace


def redact_document(document: Document, entities: List[PIIEntity]
                    ) -> Tuple[List[Redaction], str, Document, List[Dict[str, Any]]]:
    try:
        from src.redaction.redactor import redact_document as _team_redact
        result = _team_redact(document, entities)
        redactions, redacted_text, redacted_doc, trace = result
        if not isinstance(trace, list) or (trace and "tag" not in trace[0]):
            trace = _fallback_redact_document(document, entities)[3]
        return redactions, redacted_text, redacted_doc, trace
    except Exception:
        return _fallback_redact_document(document, entities)


def _fallback_report(document: Document, entities: List[PIIEntity],
                     redactions: List[Redaction], traceability_log: List[Dict[str, Any]],
                     redacted_text: str) -> Dict[str, Any]:
    counts_by_type: Dict[str, int] = {}
    tier_counts = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0}
    for e in entities:
        counts_by_type[e.entity_type] = counts_by_type.get(e.entity_type, 0) + 1
        tier = (e.metadata or {}).get("sensitivity", SENSITIVITY.get(e.entity_type, "LOW"))
        tier_counts[tier] = tier_counts.get(tier, 0) + 1
    if tier_counts.get("CRITICAL"):
        risk = "CRITICAL"
    elif tier_counts.get("HIGH"):
        risk = "HIGH"
    elif tier_counts.get("MEDIUM"):
        risk = "MEDIUM"
    elif tier_counts.get("LOW"):
        risk = "LOW"
    else:
        risk = "NONE"
    low_conf_blocks = [b.block_id for b in document.blocks
                       if (b.metadata or {}).get("low_confidence_flag")]
    audit_summary = (
        f"Artifact '{document.filename}' ({document.file_type.upper()}) processed: "
        f"{len(entities)} PII entities detected across {len(document.blocks)} structural blocks; "
        f"{len(redactions)} deterministic transparent-tag redactions applied; "
        f"{len(low_conf_blocks)} low-confidence OCR block(s) flagged for auditor review; "
        f"portfolio risk level: {risk}."
    )
    return {
        "document_id": document.document_id,
        "filename": document.filename,
        "file_type": document.file_type,
        "risk_level": risk,
        "total_pii_found": len(entities),
        "total_redactions_applied": len(redactions),
        "counts_by_type": counts_by_type,
        "counts_by_sensitivity": tier_counts,
        "low_confidence_ocr_blocks": low_conf_blocks,
        "llm_ready": True,
        "safe_text_for_llm": redacted_text,
        "traceability_log": traceability_log,
        "audit_summary": audit_summary,
    }


def generate_exposure_report(document: Document, entities: List[PIIEntity],
                             redactions: List[Redaction],
                             traceability_log: List[Dict[str, Any]],
                             redacted_text: str) -> Dict[str, Any]:
    try:
        from src.redaction.report import generate_exposure_report as _team_report
        report = _team_report(document, entities, redactions, traceability_log, redacted_text)
        if isinstance(report, dict) and "risk_level" in report:
            for key, val in _fallback_report(document, entities, redactions,
                                             traceability_log, redacted_text).items():
                report.setdefault(key, val)
            return report
    except Exception:
        pass
    return _fallback_report(document, entities, redactions, traceability_log, redacted_text)


# --------------------------------------------------------------------------- #
# 6. FAIL-SAFE NATIVE FILE EXPORTERS (DOCX / PPTX / PDF)
# --------------------------------------------------------------------------- #
def export_redacted_docx(raw_bytes: bytes, redactions: List[Redaction], entities: List[PIIEntity]) -> bytes:
    import docx as _docx_mod
    doc = _docx_mod.Document(io.BytesIO(raw_bytes))

    tag_map: Dict[str, str] = {}
    for r in redactions:
        val = (getattr(r, "original_text", "") or "").strip()
        rep = getattr(r, "replacement", "[REDACTED]")
        if val:
            tag_map[val] = rep
    for e in entities:
        val = (getattr(e, "text", "") or "").strip()
        if val and val not in tag_map:
            tag_map[val] = f"[{TAG_PREFIX.get(e.entity_type, e.entity_type.upper())}]"

    replacements = sorted(tag_map.items(), key=lambda x: len(x[0]), reverse=True)

    def _replace_in_paragraph(para):
        full_text = para.text
        for orig, rep in replacements:
            if orig in full_text:
                full_text = full_text.replace(orig, rep)
        if full_text != para.text:
            para.text = full_text

    for p in doc.paragraphs:
        _replace_in_paragraph(p)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_in_paragraph(p)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def export_redacted_pptx(raw_bytes: bytes, redactions: List[Redaction], entities: List[PIIEntity]) -> bytes:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(raw_bytes))

    tag_map: Dict[str, str] = {}
    for r in redactions:
        val = (getattr(r, "original_text", "") or "").strip()
        rep = getattr(r, "replacement", "[REDACTED]")
        if val:
            tag_map[val] = rep
    for e in entities:
        val = (getattr(e, "text", "") or "").strip()
        if val and val not in tag_map:
            tag_map[val] = f"[{TAG_PREFIX.get(e.entity_type, e.entity_type.upper())}]"

    replacements = sorted(tag_map.items(), key=lambda x: len(x[0]), reverse=True)

    def _replace_in_frame(tf):
        for para in tf.paragraphs:
            full_text = para.text
            for orig, rep in replacements:
                if orig in full_text:
                    full_text = full_text.replace(orig, rep)
            if full_text != para.text:
                para.text = full_text

    for slide in prs.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                _replace_in_frame(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        _replace_in_frame(cell.text_frame)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


def export_redacted_pdf(
    raw_bytes: bytes,
    redactions: List[Redaction],
    entities: List[PIIEntity],
    blocks: List[TextBlock],
    progress=None,
) -> bytes:
    """
    PDF redaction via app/robust_redactor.py.
    Image-only PDFs (e.g. Word "save as picture") have no text layer, so pages are OCR'd
    and the OCR word boxes are redacted (pixels under each box are blanked).
    """
    if not raw_bytes:
        raise ValueError("Original PDF bytes are missing.")
    try:
        from app.robust_redactor import redact_pdf_parallel
    except ImportError:
        try:
            from robust_redactor import redact_pdf_parallel
        except ImportError:
            from src.redaction.robust_redactor import redact_pdf_parallel

    literals: Dict[str, str] = {}
    for r in redactions:
        t = (getattr(r, "original_text", "") or "").strip().rstrip(".,;:!?)]}>(\"'")
        if len(t) >= 3:
            etype = getattr(r, "entity_type", "")
            if etype in ("PERSON", "PERSON_NAME") and not is_valid_person_name(t):
                continue
            literals[t] = f"[{TAG_PREFIX.get(etype, str(etype).upper())}]"
    for e in entities:
        t = (getattr(e, "text", "") or "").strip().rstrip(".,;:!?)]}>(\"'")
        if len(t) >= 3:
            etype = getattr(e, "entity_type", "")
            if etype in ("PERSON", "PERSON_NAME") and not is_valid_person_name(t):
                continue
            literals.setdefault(t, f"[{TAG_PREFIX.get(etype, etype.upper())}]")

    out, stats = redact_pdf_parallel(raw_bytes, literals=literals, progress=progress)
    if stats["leaks"]:
        raise RuntimeError(f"Zero-leakage check failed: {len(stats['leaks'])} residual match(es), "
                           f"e.g. {stats['leaks'][:3]}")
    return out


def generate_native_export(filename: str, raw_bytes: bytes, redactions: List[Redaction],
                           entities: List[PIIEntity], blocks: List[TextBlock]) -> Tuple[Optional[bytes], str, str]:
    ext = Path(filename).suffix.lower()
    if ext == ".docx":
        return export_redacted_docx(raw_bytes, redactions, entities), f"{Path(filename).stem}_redacted.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif ext == ".pptx":
        return export_redacted_pptx(raw_bytes, redactions, entities), f"{Path(filename).stem}_redacted.pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    elif ext == ".pdf":
        return export_redacted_pdf(raw_bytes, redactions, entities, blocks), f"{Path(filename).stem}_redacted.pdf", "application/pdf"
    return None, "", ""


# --------------------------------------------------------------------------- #
# 7. METRIC HELPERS & ZERO-LEAKAGE CHECK
# --------------------------------------------------------------------------- #
def risk_badge(risk: str) -> str:
    cls = {"CRITICAL": "risk-critical", "HIGH": "risk-high", "MEDIUM": "risk-medium",
           "LOW": "risk-low", "NONE": "risk-none"}.get(risk, "risk-none")
    return f'<span class="pill {cls}">{risk}</span>'


def verify_zero_leakage(entities: List[PIIEntity], redacted_text: str) -> Tuple[bool, List[str]]:
    leaked = []
    for e in entities:
        t = (e.text or "").strip()
        if t and t in redacted_text:
            leaked.append(t)
    residual_patterns = re.findall(
        r"\b\d{3}-\d{2}-\d{4}\b|[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", redacted_text or "")
    return (len(leaked) == 0 and not residual_patterns), leaked + residual_patterns


def records_to_csv(records: List[Dict[str, Any]]) -> str:
    if not records:
        return ""
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=list(records[0].keys()))
    writer.writeheader()
    writer.writerows(records)
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# 8. PIPELINE ORCHESTRATION
# --------------------------------------------------------------------------- #
def run_pipeline(path: Union[str, Path], threshold: float) -> Dict[str, Any]:
    filename = Path(path).name
    with st.status(f"Processing artifact: {filename}", expanded=True) as status:
        st.write("1/5 Ingestion — routing artifact by file format...")
        document = route_and_extract(path)
        st.write(f"2/5 Extraction — {len(document.blocks)} structural blocks extracted.")
        entities = detect_pii(document, threshold)
        st.write(f"3/5 Detection — {len(entities)} sensitive entities identified.")
        redactions, redacted_text, redacted_doc, trace = redact_document(document, entities)
        st.write(f"4/5 Redaction — {len(redactions)} transparent tags applied.")
        report = generate_exposure_report(document, entities, redactions, trace, redacted_text)
        st.write("5/5 Audit Report — Exposure and compliance reports ready.")
        status.update(label=f"Completed: {filename} (Risk: {report.get('risk_level', 'NONE')})", state="complete")

    return {
        "document": document, "entities": entities, "redactions": redactions,
        "redacted_text": redacted_text, "redacted_doc": redacted_doc,
        "trace": trace, "report": report,
    }


def find_sample_files() -> List[Path]:
    candidates: List[Path] = []
    for sub in ("tests/data", "test_data", "samples", "data", "examples"):
        d = PROJECT_ROOT / sub
        if d.is_dir():
            for ext in ("*.pdf", "*.docx", "*.pptx"):
                candidates.extend(sorted(d.glob(ext)))
    return candidates


def process_uploaded_files(uploaded_files, threshold: float) -> None:
    for up in uploaded_files:
        suffix = Path(up.name).suffix.lower()
        if suffix not in (".docx", ".pdf", ".pptx"):
            st.warning(f"Skipped {up.name}: unsupported type {suffix}")
            continue

        file_bytes = up.getvalue()
        st.session_state.raw_files[up.name] = file_bytes

        if up.name in st.session_state.results:
            res_doc = st.session_state.results[up.name].get("document")
            if res_doc:
                st.session_state.raw_files[res_doc.filename] = file_bytes
            continue

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
        try:
            tmp.write(file_bytes)
            tmp.close()
            result = run_pipeline(tmp.name, threshold)
            st.session_state.results[up.name] = result

            if result.get("document"):
                st.session_state.raw_files[result["document"].filename] = file_bytes
        except Exception as exc:
            st.error(f"Failed to process {up.name}: {exc}")
        finally:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass


# --------------------------------------------------------------------------- #
# 9. UI & STYLING
# --------------------------------------------------------------------------- #
def inject_css() -> None:
    st.markdown("""
<style>
:root {
  --bg: #0b1220; --panel: rgba(20, 30, 48, 0.72); --border: rgba(120, 144, 180, 0.22);
  --txt: #e2e8f0; --mut: #94a3b8; --accent: #38bdf8; --accent2: #818cf8;
}
.stApp { background: radial-gradient(1200px 600px at 10% -10%, #16233d 0%, var(--bg) 55%); color: var(--txt); }
[data-testid="stSidebar"] { background: linear-gradient(180deg, #0e1729 0%, #0b1220 100%); }
.brand-banner {
  display: flex; align-items: center; justify-content: space-between;
  padding: 18px 26px; border: 1px solid var(--border); border-radius: 16px;
  background: linear-gradient(120deg, rgba(56,189,248,.10), rgba(129,140,248,.10)), var(--panel);
  backdrop-filter: blur(10px); margin-bottom: 22px;
}
.brand-left { font-weight: 800; font-size: 22px; letter-spacing: 2px; color: #38bdf8; }
.brand-mid { text-align: center; }
.brand-mid .title { font-size: 24px; font-weight: 700; color: #f8fafc; }
.brand-mid .sub { color: var(--mut); font-size: 13px; margin-top: 2px; }
.brand-right { font-weight: 800; font-size: 22px; letter-spacing: 2px; color: #f8fafc; }
.kpi-card {
  border: 1px solid var(--border); border-radius: 14px; padding: 16px 18px;
  background: var(--panel); backdrop-filter: blur(8px); min-height: 108px;
}
.kpi-card .kpi-label { color: var(--mut); font-size: 12px; text-transform: uppercase; letter-spacing: 1px; }
.kpi-card .kpi-value { font-size: 28px; font-weight: 800; margin-top: 4px; color: #f8fafc; }
.kpi-card .kpi-foot { font-size: 12px; margin-top: 4px; }
.pill { display: inline-block; padding: 3px 12px; border-radius: 999px; font-size: 12px; font-weight: 700; }
.risk-critical { background: rgba(239,68,68,.18); color: #f87171; border: 1px solid rgba(239,68,68,.5); }
.risk-high { background: rgba(249,115,22,.18); color: #fb923c; border: 1px solid rgba(249,115,22,.5); }
.risk-medium { background: rgba(245,158,11,.18); color: #fbbf24; border: 1px solid rgba(245,158,11,.5); }
.risk-low, .risk-none { background: rgba(16,185,129,.16); color: #34d399; border: 1px solid rgba(16,185,129,.5); }
.shield {
  display: inline-flex; align-items: center; gap: 8px; padding: 10px 18px; border-radius: 12px;
  background: rgba(16,185,129,.12); border: 1px solid rgba(16,185,129,.45); color: #34d399; font-weight: 700;
}
.shield-bad { background: rgba(239,68,68,.12); border-color: rgba(239,68,68,.5); color: #f87171; }
.panel { border: 1px solid var(--border); border-radius: 14px; background: var(--panel); padding: 16px; backdrop-filter: blur(8px); }
.doc-pane { border: 1px solid var(--border); border-radius: 12px; padding: 14px; background: rgba(11,18,32,.6); max-height: 560px; overflow-y: auto; font-size: 14px; line-height: 1.7; }
.doc-pane h4 { color: var(--accent); margin: 0 0 6px 0; font-size: 13px; text-transform: uppercase; letter-spacing: 1px; }
.block-label { color: var(--mut); font-size: 11px; margin: 14px 0 2px 0; font-family: monospace; }
.heading-block { font-weight: 700; font-size: 16px; color: #f1f5f9; }
pre.safe-text { white-space: pre-wrap; word-wrap: break-word; font-size: 13px; line-height: 1.7; color: #d1fae5; background: rgba(6,26,20,.5); padding: 12px; border-radius: 10px; }
mark { padding: 1px 5px; border-radius: 6px; font-weight: 700; color: #0b1220; }
mark.mk-person { background: #7dd3fc; } mark.mk-employee_id { background: #c4b5fd; }
mark.mk-email_address { background: #67e8f9; } mark.mk-phone_number { background: #5eead4; }
mark.mk-us_ssn { background: #fca5a5; } mark.mk-government_id { background: #f87171; }
mark.mk-credit_card { background: #fdba74; } mark.mk-account_number { background: #fcd34d; }
mark.mk-date_of_birth { background: #fde68a; } mark.mk-address { background: #86efac; }
.bar-row { display: flex; align-items: center; gap: 10px; margin: 6px 0; }
.bar-row .bar-lbl { width: 170px; font-size: 13px; color: var(--txt); font-family: monospace; }
.bar-row .bar-track { flex: 1; background: rgba(148,163,184,.15); border-radius: 6px; height: 18px; }
.bar-row .bar-fill { height: 18px; border-radius: 6px; background: linear-gradient(90deg, #38bdf8, #818cf8); }
.bar-row .bar-val { width: 44px; text-align: right; font-weight: 700; }
.stage-line { font-size: 13px; margin: 4px 0; color: var(--mut); }
.stage-done { color: #34d399; font-weight: 700; }
.risk-banner { padding: 14px 18px; border-radius: 12px; font-weight: 600; margin: 10px 0; border: 1px solid var(--border); }
.justification { font-size: 13px; color: var(--mut); margin-top: 6px; }
.export-card { border: 1px solid rgba(56,189,248,.3); border-radius: 12px; padding: 14px; background: rgba(56,189,248,.05); margin-top: 14px; }
footer { visibility: hidden; }
</style>
""", unsafe_allow_html=True)


def render_header() -> None:
    st.markdown("""
<div class="brand-banner">
  <div class="brand-mid" style="text-align:left">
    <div class="title">Cadence PII Shield</div>
    <div class="sub">Detect &bull; Redact &bull; Export &bull; Audit</div>
  </div>
</div>
""", unsafe_allow_html=True)


def render_sidebar() -> Tuple[float, List[str]]:
    st.sidebar.markdown("### ⚙️ Controls")
    threshold = st.sidebar.slider("Detection Confidence Threshold", 0.50, 1.00, 0.75, 0.05)
    type_filter = st.sidebar.multiselect(
        "Redact only these categories (default: all)",
        options=ALL_TYPES,
        format_func=lambda t: f"{TYPE_SHORT.get(t, t)} ({SENSITIVITY.get(t, 'LOW')})",
        default=[])
    st.sidebar.markdown("---")
    uploaded = st.sidebar.file_uploader(
        "Upload (.pdf, .docx, .pptx)",
        type=["docx", "pdf", "pptx"], accept_multiple_files=True)
    if uploaded:
        process_uploaded_files(uploaded, threshold)
    if st.sidebar.button("🗑 Reset session"):
        st.session_state.results = {}
        st.session_state.raw_files = {}
        st.session_state.exports = {}
        st.rerun()
    return threshold, type_filter


def apply_type_filter(results: Dict[str, Any], type_filter: List[str]) -> Dict[str, Any]:
    if not type_filter:
        return results
    keep = set()
    for t in type_filter:
        keep.add(t)
        for alias, canon in CANONICAL_TYPE_MAP.items():
            if canon == t or alias == t:
                keep.add(alias)
                keep.add(canon)

    out = dict(results)
    out["entities"] = [e for e in results["entities"] if e.entity_type in keep or CANONICAL_TYPE_MAP.get(e.entity_type) in keep]
    out["redactions"] = [r for r in results["redactions"] if r.entity_type in keep or CANONICAL_TYPE_MAP.get(r.entity_type) in keep]
    out["trace"] = [t for t in results["trace"] if t.get("pii_type") in keep or CANONICAL_TYPE_MAP.get(t.get("pii_type")) in keep]
    report = dict(results["report"])
    counts = {k: v for k, v in report.get("counts_by_type", {}).items() if k in keep or CANONICAL_TYPE_MAP.get(k) in keep}
    report["counts_by_type"] = counts
    report["total_pii_found"] = sum(counts.values())
    report["total_redactions_applied"] = len(out["redactions"])
    tiers = {"CRITICAL": 0, "HIGH": 0, "MEDIUM": 0, "LOW": 0}
    for t in out["trace"]:
        tiers[t.get("sensitivity", "LOW")] = tiers.get(t.get("sensitivity", "LOW"), 0) + 1
    report["counts_by_sensitivity"] = tiers
    report["risk_level"] = ("CRITICAL" if tiers["CRITICAL"] else
                            "HIGH" if tiers["HIGH"] else
                            "MEDIUM" if tiers["MEDIUM"] else
                            "LOW" if tiers["LOW"] else "NONE")
    out["report"] = report
    return out


def render_portfolio_summary() -> None:
    results = st.session_state.results
    if not results:
        st.info("Upload a document in the sidebar to start.")
        return
    total_pii = sum(r["report"].get("total_pii_found", 0) for r in results.values())
    highest = "NONE"
    for r in results.values():
        lvl = r["report"].get("risk_level", "NONE")
        if RISK_ORDER.get(lvl, 0) > RISK_ORDER.get(highest, 0):
            highest = lvl
    cards = [
        ("Documents", str(len(results))),
        ("PII Entities Found", f"{total_pii:,}"),
        ("Highest Risk", f'<span class="pill risk-{highest.lower()}">{highest}</span>'),
    ]
    for col, (label, value) in zip(st.columns(3), cards):
        col.markdown(f'<div class="kpi-card"><div class="kpi-label">{label}</div>'
                     f'<div class="kpi-value">{value}</div></div>', unsafe_allow_html=True)
    st.markdown("")


def html_bar_chart(counts: Dict[str, int], unit: str = "") -> str:
    if not counts:
        return '<div class="panel" style="color:#94a3b8">No items recorded.</div>'
    maxv = max(counts.values()) or 1
    rows = []
    for label, val in sorted(counts.items(), key=lambda kv: -kv[1]):
        pct = max(4, round(100 * val / maxv))
        rows.append(f'<div class="bar-row"><div class="bar-lbl">{html.escape(str(label))}</div>'
                    f'<div class="bar-track"><div class="bar-fill" style="width:{pct}%"></div></div>'
                    f'<div class="bar-val">{val}{unit}</div></div>')
    return '<div class="panel">' + "".join(rows) + "</div>"


def render_comparison_tab(res: Dict[str, Any], selected_key: str = "") -> None:
    doc: Document = res["document"]
    entities: List[PIIEntity] = res["entities"]
    per_block: Dict[str, List[PIIEntity]] = {}
    for e in entities:
        per_block.setdefault(e.block_id or "", []).append(e)

    def block_html(block: TextBlock, redacted: bool) -> str:
        text = block.text or ""
        if redacted:
            return f'<pre class="safe-text" style="margin:6px 0 12px 0">{html.escape(text)}</pre>'
        ents = sorted(per_block.get(block.block_id, []), key=lambda e: e.start)
        parts, cursor = [], 0
        for e in ents:
            if e.start < cursor:
                continue
            parts.append(html.escape(text[cursor:e.start]))
            cls = "mk-" + e.entity_type.lower()
            parts.append(f'<mark class="{cls}" title="{e.entity_type} (conf {e.confidence:.2f})">{html.escape(text[e.start:e.end])}</mark>')
            cursor = e.end
        parts.append(html.escape(text[cursor:]))
        body = "".join(parts).replace("\n", "<br>")
        cls = "heading-block" if (block.metadata or {}).get("is_heading") else ""
        loc = f"p.{block.page_number}" if block.page_number else "n/a"
        return (f'<div class="block-label">[{block.block_id} | {loc} | '
                f'{(block.metadata or {}).get("block_type", "paragraph")}]</div>'
                f'<div class="{cls}">{body}</div>')

    raw_html = "".join(block_html(b, False) for b in doc.blocks) or "<i>No text extracted.</i>"
    safe_html = "".join(block_html(b, True) for b in res["redacted_doc"].blocks) or "<i>No text extracted.</i>"

    left, right = st.columns(2)
    with left:
        st.markdown("#### Raw Evidence (Detected PII)")
        st.markdown(f'<div class="doc-pane"><h4>Extracted Source Elements</h4>{raw_html}</div>', unsafe_allow_html=True)
    with right:
        st.markdown("#### Redacted Output (Sanitized)")
        ok, leaked = verify_zero_leakage(res["entities"], res["redacted_text"])
        if ok:
            st.markdown('<div class="shield">🛡 ZERO PII LEAKAGE VERIFIED</div>', unsafe_allow_html=True)
        else:
            st.markdown(f'<div class="shield shield-bad">⛔ Potential Residual PII: {len(leaked)} token(s)</div>', unsafe_allow_html=True)
        st.markdown(f'<div class="doc-pane"><h4>Sanitized Payload</h4>{safe_html}</div>', unsafe_allow_html=True)

    render_export_section(res, selected_key)


def _find_raw_bytes(doc: Document, selected_key: str) -> Optional[bytes]:
    raw = (st.session_state.raw_files.get(selected_key)
           or st.session_state.raw_files.get(doc.filename))
    if not raw and len(st.session_state.raw_files) == 1:
        raw = next(iter(st.session_state.raw_files.values()))
    if not raw:
        for smp in find_sample_files():
            if smp.name in (selected_key, doc.filename) and smp.exists():
                raw = smp.read_bytes()
                break
    return raw


def _export_cache_key(raw: bytes, res: Dict[str, Any]) -> str:
    h = hashlib.sha256(raw)
    for t in sorted({(e.entity_type, e.text) for e in res["entities"]}):
        h.update(repr(t).encode("utf-8", "ignore"))
    return h.hexdigest()


def render_export_section(res: Dict[str, Any], selected_key: str) -> None:
    doc: Document = res["document"]
    name = doc.filename or selected_key
    stem, ext = Path(name).stem, Path(name).suffix.lower()
    cache = st.session_state.setdefault("exports", {})
    raw = _find_raw_bytes(doc, selected_key)

    st.markdown("##### 📥 Downloads")
    c_native, c_txt = st.columns(2)

    with c_native:
        if not raw:
            st.info("Original file bytes missing. Re-upload the document to enable export.")
        elif ext == ".pdf":
            key = _export_cache_key(raw, res)
            if key not in cache:
                st.caption("Scanned / image-only PDFs are OCR'd page by page, which can take a minute or two.")
                if st.button("⚙ Generate redacted PDF", key=f"gen_{key[:16]}", use_container_width=True):
                    bar = st.progress(0.0, text="Starting OCR + redaction...")
                    try:
                        cache[key] = export_redacted_pdf(
                            raw, res["redactions"], res["entities"], doc.blocks,
                            progress=lambda d, n: bar.progress(d / n, text=f"Redacted {d}/{n} pages"))
                        failed = None
                    except Exception as exc:
                        failed = exc
                    bar.empty()
                    if failed is not None:
                        st.error(f"PDF export failed: {failed}")
                    else:
                        st.rerun()
            if key in cache:
                st.download_button("⬇ Download Redacted PDF", data=cache[key],
                                   file_name=f"{stem}_redacted.pdf", mime="application/pdf",
                                   use_container_width=True)
        else:
            data, out_name, mime = generate_native_export(
                name, raw, res["redactions"], res["entities"], doc.blocks)
            if data:
                st.download_button(f"⬇ Download Redacted {ext.upper().lstrip('.')}", data=data,
                                   file_name=out_name, mime=mime, use_container_width=True)

    with c_txt:
        st.download_button("⬇ Download Safe Text (.txt)", res["redacted_text"],
                           file_name=f"{stem}_sanitized.txt", mime="text/plain",
                           use_container_width=True)


def render_audit_tab(res: Dict[str, Any]) -> None:
    report = res["report"]
    risk = report.get("risk_level", "NONE")
    counts = report.get("counts_by_type", {})
    tiers = report.get("counts_by_sensitivity", {})
    justifications = {
        "CRITICAL": "Critical regulatory identifiers or financial credentials were detected. Strict masking required.",
        "HIGH": "Employee IDs or sensitive demographic attributes detected. Identity protection measures active.",
        "MEDIUM": "Direct contact details present. Phishing/harvesting risk mitigated.",
        "LOW": "Person names detected. Basic anonymization applied.",
        "NONE": "No sensitive PII detected above threshold.",
    }
    st.markdown(f"**Document Exposure Posture:** {risk_badge(risk)}")
    st.markdown(f'<div class="risk-banner panel">{html.escape(justifications.get(risk, ""))}'
                f'<div class="justification">{html.escape(report.get("audit_summary", ""))}</div></div>',
                unsafe_allow_html=True)

    left, right = st.columns(2)
    with left:
        st.markdown("**Entity Counts by Category**")
        st.markdown(html_bar_chart(counts), unsafe_allow_html=True)
    with right:
        st.markdown("**Entity Counts by Risk Level**")
        st.markdown(html_bar_chart(tiers), unsafe_allow_html=True)

    c1, c2, c3 = st.columns(3)
    c1.metric("Total Identified Entities", report.get("total_pii_found", 0))
    c2.metric("Redactions Applied", report.get("total_redactions_applied", 0))
    c3.metric("Flagged OCR Regions", len(report.get("low_confidence_ocr_blocks", [])))

    st.markdown("##### Download Compliance Audit Reports")
    d1, d2 = st.columns(2)
    report_json = json.dumps(report, indent=2, default=str)
    d1.download_button("⬇ Download Audit Report (.json)", report_json,
                       file_name=f"{Path(report.get('filename', 'artifact')).stem}_audit.json",
                       mime="application/json", use_container_width=True)
    flat_records = [{
        "filename": report.get("filename"), "risk_level": risk,
        "entity_type": k, "count": v, "sensitivity": SENSITIVITY.get(k, "LOW"),
    } for k, v in counts.items()]
    d2.download_button("⬇ Download Audit Report (.csv)", records_to_csv(flat_records),
                       file_name=f"{Path(report.get('filename', 'artifact')).stem}_audit.csv",
                       mime="text/csv", use_container_width=True)


def render_traceability_tab(res: Dict[str, Any]) -> None:
    trace: List[Dict[str, Any]] = list(res["trace"])
    st.markdown("Audit Log — Real-time mapping of applied tags to original text spans.")
    col_a, col_b = st.columns([3, 1])
    query = col_a.text_input("🔍 Search Traceability Log", "")
    reveal = col_b.toggle("Reveal Original Values", value=False)

    rows = trace
    if query:
        q = query.lower()
        rows = [t for t in rows if q in json.dumps(t, default=str).lower()]
    display_rows = []
    for t in rows:
        r = dict(t)
        if not reveal:
            r["original_text"] = "•" * min(12, max(6, len(str(r.get("original_text", "")))))
        display_rows.append(r)

    if _HAVE_PANDAS:
        st.dataframe(pd.DataFrame(display_rows), use_container_width=True, hide_index=True)
    else:
        st.table(display_rows)
    st.caption(f"{len(display_rows)} log records")

    st.download_button("⬇ Export Traceability (.csv)", records_to_csv(display_rows),
                       file_name="traceability_log.csv", mime="text/csv")


def render_artifact_workspace(type_filter: List[str]) -> None:
    results = st.session_state.results
    if not results:
        return
    names = sorted(results.keys())
    selected = st.selectbox("Select Active Document", names, index=len(names) - 1)
    res = apply_type_filter(results[selected], type_filter)

    head1, head2, head3 = st.columns([3, 2, 2])
    head1.markdown(f"**Document:** `{selected}`")
    head2.markdown(f"**Risk Level:** {risk_badge(res['report'].get('risk_level', 'NONE'))}", unsafe_allow_html=True)
    head3.markdown(f"**PII / Redactions:** {res['report'].get('total_pii_found', 0)} / {res['report'].get('total_redactions_applied', 0)}")

    tab1, tab2, tab3 = st.tabs(["Redaction & Downloads", "Audit Report", "Traceability"])
    with tab1:
        render_comparison_tab(res, selected)
    with tab2:
        render_audit_tab(res)
    with tab3:
        render_traceability_tab(res)


# --------------------------------------------------------------------------- #
# 10. ENTRYPOINT
# --------------------------------------------------------------------------- #
def main() -> None:
    st.set_page_config(page_title="Cadence | PII Shield", page_icon="🛡", layout="wide", initial_sidebar_state="expanded")
    if "results" not in st.session_state:
        st.session_state.results = {}
    if "raw_files" not in st.session_state:
        st.session_state.raw_files = {}
    if "exports" not in st.session_state:
        st.session_state.exports = {}

    inject_css()
    render_header()
    threshold, type_filter = render_sidebar()
    render_portfolio_summary()
    render_artifact_workspace(type_filter)


if __name__ == "__main__":
    main()