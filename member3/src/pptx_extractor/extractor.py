"""Main PPTX Extraction Engine for Member 3.
Cadence Financial Group - Case Study 2: Text Analysis & PII Detection.

Orchestrates slide traversal, shape parsing, table extraction, chart reading,
speaker notes, image OCR, deterministic ordering, and context grouping.
"""

from __future__ import annotations
import os
from pathlib import Path
from typing import List, Dict, Any, Optional
import pptx
from pptx.enum.shapes import MSO_SHAPE_TYPE

from .schema import ExtractedBlock, PPTXExtractionResult, SlideSummary
from .shape_extractor import ShapeExtractor
from .table_extractor import TableExtractor
from .chart_extractor import ChartExtractor
from .notes_extractor import NotesExtractor
from .image_ocr import ImageOCRExtractor, BaseOCRAdapter
from .ordering import DeterministicOrderingEngine
from .context_grouping import ContextGroupingEngine


class PPTXExtractor:
    """Master extraction pipeline for PowerPoint presentations."""

    def __init__(self, ocr_adapter: Optional[BaseOCRAdapter] = None):
        self.ocr_adapter = ocr_adapter
        self.ordering_engine = DeterministicOrderingEngine()
        self.context_grouping_engine = ContextGroupingEngine()

    def extract(self, pptx_path: str | Path) -> PPTXExtractionResult:
        """Extract all content, structures, tables, charts, notes, and images from a PPTX file."""
        file_path = Path(pptx_path)
        if not file_path.exists():
            raise FileNotFoundError(f"PPTX file not found: {pptx_path}")

        file_name = file_path.name
        prs = pptx.Presentation(str(file_path))

        shape_extractor = ShapeExtractor(file_name)
        table_extractor = TableExtractor(file_name)
        chart_extractor = ChartExtractor(file_name)
        notes_extractor = NotesExtractor(file_name)
        image_extractor = ImageOCRExtractor(file_name, ocr_adapter=self.ocr_adapter)

        all_extracted_blocks: List[ExtractedBlock] = []
        slide_summaries: List[SlideSummary] = []
        global_order_counter = 0

        for slide_idx, slide in enumerate(prs.slides, start=1):
            slide_id = getattr(slide, "slide_id", None)
            slide_blocks: List[ExtractedBlock] = []
            table_count = 0
            chart_count = 0
            image_count = 0
            titles_found: List[str] = []

            # 1. Process all visual shapes on the slide
            for shape in slide.shapes:
                # A. Table Shape
                if getattr(shape, "has_table", False):
                    table_count += 1
                    t_blocks, global_order_counter = table_extractor.extract_from_table(
                        shape=shape,
                        slide_num=slide_idx,
                        slide_id=slide_id,
                        table_index=table_count,
                        order_counter=global_order_counter,
                    )
                    slide_blocks.extend(t_blocks)

                # B. Chart Shape
                elif getattr(shape, "has_chart", False):
                    chart_count += 1
                    c_blocks, global_order_counter = chart_extractor.extract_from_chart(
                        shape=shape,
                        slide_num=slide_idx,
                        slide_id=slide_id,
                        chart_index=chart_count,
                        order_counter=global_order_counter,
                    )
                    slide_blocks.extend(c_blocks)

                # C. Picture Shape (Image)
                elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
                    image_count += 1
                    img_block, global_order_counter = image_extractor.extract_from_picture(
                        shape=shape,
                        slide_num=slide_idx,
                        slide_id=slide_id,
                        image_index=image_count,
                        order_counter=global_order_counter,
                    )
                    if img_block:
                        slide_blocks.append(img_block)

                # D. Text Box, Auto Shape, Title, or Grouped Shapes
                else:
                    s_blocks, global_order_counter = shape_extractor.extract_from_shape(
                        shape=shape,
                        slide_num=slide_idx,
                        slide_id=slide_id,
                        order_counter=global_order_counter,
                    )
                    for sb in s_blocks:
                        if sb.block_type == "title" and sb.text:
                            titles_found.append(sb.text.split("\n")[0])
                    slide_blocks.extend(s_blocks)

            # 2. Process Speaker Notes
            notes_block, global_order_counter = notes_extractor.extract_from_slide(
                slide=slide,
                slide_num=slide_idx,
                slide_id=slide_id,
                order_counter=global_order_counter,
            )
            has_notes = False
            if notes_block:
                has_notes = True
                slide_blocks.append(notes_block)

            # 3. Create Slide Summary
            summary = SlideSummary(
                slide_number=slide_idx,
                slide_id=slide_id,
                total_blocks=len(slide_blocks),
                titles=titles_found,
                has_tables=(table_count > 0),
                has_charts=(chart_count > 0),
                has_images=(image_count > 0),
                has_notes=has_notes,
                context_groups=[],
            )
            slide_summaries.append(summary)
            all_extracted_blocks.extend(slide_blocks)

        # 4. Apply Deterministic 2D Reading Order
        ordered_blocks = self.ordering_engine.order_blocks(all_extracted_blocks)

        # 5. Apply Context Grouping across ordered blocks
        context_groups_map = self.context_grouping_engine.process_all_blocks(ordered_blocks)

        # Update slide summary context group IDs
        for summary in slide_summaries:
            prefix = f"slide_{summary.slide_number}_group_"
            summary.context_groups = [
                gid for gid in context_groups_map.keys() if gid.startswith(prefix)
            ]

        # 6. Assemble Full Presentation Extraction Result
        ocr_status = image_extractor.get_status()
        extraction_result = PPTXExtractionResult(
            file_name=file_name,
            total_slides=len(prs.slides),
            extracted_slides=len(prs.slides),
            total_blocks=len(ordered_blocks),
            blocks=ordered_blocks,
            slide_summaries=slide_summaries,
            context_groups=context_groups_map,
            metadata={
                "ocr_status": ocr_status,
                "ordering_algorithm": "deterministic_2d_spatial_banding",
                "context_clustering": "spatial_proximity_and_hierarchy",
                "team_role": "Member 3 - PowerPoint Extraction & Structure Retention",
            },
        )

        return extraction_result
