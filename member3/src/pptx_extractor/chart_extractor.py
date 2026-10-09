"""Chart Extractor for Member 3 PPTX Engine.
Cadence Financial Group - Case Study 2: Text Analysis & PII Detection.

Extracts titles, axis labels, categories, and series data from charts
using python-pptx, with documented fallbacks for unsupported chart properties.
"""

from __future__ import annotations
from typing import List, Tuple, Any, Dict
import pptx

from .schema import ExtractedBlock, Position
from .shape_extractor import extract_shape_position


class ChartExtractor:
    """Extracts text, series, and category labels from embedded charts."""

    def __init__(self, file_name: str):
        self.file_name = file_name

    def extract_from_chart(
        self,
        shape: Any,
        slide_num: int,
        slide_id: int | None,
        chart_index: int,
        order_counter: int,
    ) -> Tuple[List[ExtractedBlock], int]:
        """Extract all accessible text content from a chart shape.
        Returns a tuple of (blocks, updated_order_counter).
        """
        blocks: List[ExtractedBlock] = []
        chart = shape.chart
        pos = extract_shape_position(shape)
        chart_ctx_id = f"s{slide_num}_chart_{chart_index}"

        chart_title = None
        limitations: List[str] = []

        # 1. Chart Title
        try:
            if chart.has_title and chart.chart_title:
                chart_title = chart.chart_title.text_frame.text.strip()
        except Exception as e:
            limitations.append(f"Chart title extraction limitation: {str(e)}")

        # 2. Categories and Series
        categories: List[str] = []
        series_names: List[str] = []
        series_data_summary: List[Dict[str, Any]] = []

        try:
            for plot in chart.plots:
                # Categories
                try:
                    if hasattr(plot, "categories") and plot.categories:
                        for cat in plot.categories:
                            cat_label = getattr(cat, "label", str(cat)).strip()
                            if cat_label and cat_label not in categories:
                                categories.append(cat_label)
                except Exception as ce:
                    limitations.append(f"Categories extraction limitation: {str(ce)}")

                # Series
                try:
                    for series in plot.series:
                        s_name = getattr(series, "name", "").strip()
                        if s_name:
                            series_names.append(s_name)
                        try:
                            s_vals = list(getattr(series, "values", []))
                            series_data_summary.append({
                                "series_name": s_name,
                                "values_count": len(s_vals),
                            })
                        except Exception:
                            pass
                except Exception as se:
                    limitations.append(f"Series extraction limitation: {str(se)}")
        except Exception as pe:
            limitations.append(f"Plot extraction limitation: {str(pe)}")

        # 3. Axis Titles
        axis_titles: List[str] = []
        try:
            cat_axis = getattr(chart, "category_axis", None)
            if cat_axis and getattr(cat_axis, "has_title", False) and cat_axis.axis_title:
                ax_text = cat_axis.axis_title.text_frame.text.strip()
                if ax_text:
                    axis_titles.append(f"Category Axis: {ax_text}")
        except Exception:
            pass

        try:
            val_axis = getattr(chart, "value_axis", None)
            if val_axis and getattr(val_axis, "has_title", False) and val_axis.axis_title:
                ax_text = val_axis.axis_title.text_frame.text.strip()
                if ax_text:
                    axis_titles.append(f"Value Axis: {ax_text}")
        except Exception:
            pass

        # Assemble extracted text lines
        text_lines: List[str] = []
        if chart_title:
            text_lines.append(f"Chart Title: {chart_title}")
        if categories:
            text_lines.append(f"Categories: {', '.join(categories)}")
        if series_names:
            text_lines.append(f"Series: {', '.join(series_names)}")
        if axis_titles:
            text_lines.extend(axis_titles)

        composite_chart_text = "\n".join(text_lines)
        if not composite_chart_text:
            composite_chart_text = f"[Chart {chart_index}: Graphical representation with no textual data labels accessible]"

        order_counter += 1
        block_id = f"s{slide_num}_chart_{chart_index}"

        block = ExtractedBlock(
            file=self.file_name,
            slide=slide_num,
            section=None,
            block_id=block_id,
            block_type="chart_text",
            text=composite_chart_text,
            position=pos,
            order=order_counter,
            context_group_id=chart_ctx_id,
            source_type="chart",
            slide_id=slide_id,
            shape_id=shape.shape_id,
            confidence=1.0,
            chart_metadata={
                "chart_index": chart_index,
                "chart_title": chart_title,
                "categories": categories,
                "series_names": series_names,
                "axis_titles": axis_titles,
                "limitations": limitations if limitations else None,
            },
        )
        blocks.append(block)

        return blocks, order_counter
