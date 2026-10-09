"""Speaker Notes Extractor for Member 3 PPTX Engine.
Cadence Financial Group - Case Study 2: Text Analysis & PII Detection.

Extracts speaker notes while preserving slide provenance and keeping notes
isolated from main visual slide content.
"""

from __future__ import annotations
from typing import Optional, List, Tuple, Any
import pptx

from .schema import ExtractedBlock, Position


class NotesExtractor:
    """Extracts speaker notes attached to individual slides."""

    def __init__(self, file_name: str):
        self.file_name = file_name

    def extract_from_slide(
        self,
        slide: Any,
        slide_num: int,
        slide_id: Optional[int],
        order_counter: int,
    ) -> Tuple[Optional[ExtractedBlock], int]:
        """Extract speaker notes if present on the slide.
        Returns a tuple of (optional_block, updated_order_counter).
        """
        if not getattr(slide, "has_notes_slide", False):
            return None, order_counter

        notes_slide = slide.notes_slide
        if not hasattr(notes_slide, "notes_text_frame"):
            return None, order_counter

        notes_tf = notes_slide.notes_text_frame
        raw_text = notes_tf.text.strip()
        if not raw_text:
            return None, order_counter

        # Extract paragraphs
        paragraphs = [p.text.strip() for p in notes_tf.paragraphs if p.text.strip()]

        order_counter += 1
        block_id = f"s{slide_num}_notes"
        ctx_group_id = f"s{slide_num}_notes"

        # Position for speaker notes is represented off-slide or standard footer metadata
        pos = Position(left=0.0, top=1000.0, width=800.0, height=200.0)

        block = ExtractedBlock(
            file=self.file_name,
            slide=slide_num,
            section=None,
            block_id=block_id,
            block_type="speaker_notes",
            text=raw_text,
            position=pos,
            order=order_counter,
            context_group_id=ctx_group_id,
            source_type="speaker_notes",
            slide_id=slide_id,
            confidence=1.0,
            paragraphs=paragraphs,
        )

        return block, order_counter
